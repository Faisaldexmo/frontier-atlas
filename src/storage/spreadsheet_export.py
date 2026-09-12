import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


logger = logging.getLogger(__name__)

OUTPUT_DIR = Path("data/output")
EXCEL_FILE = OUTPUT_DIR / "Frontier_Atlas.xlsx"


def load_json(filename: str) -> list[dict]:
    """Load a JSON list from the output directory."""

    path = OUTPUT_DIR / filename

    if not path.exists():
        raise FileNotFoundError(
            f"Required file not found: {path}"
        )

    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    if not isinstance(data, list):
        raise ValueError(
            f"{filename} must contain a JSON list."
        )

    return data


def write_sheet(
    workbook: Workbook,
    sheet_name: str,
    records: list[dict],
    columns: list[str],
) -> None:
    """Create and populate one worksheet."""

    worksheet = workbook.create_sheet(
        title=sheet_name
    )

    # Header row
    for column_index, column_name in enumerate(
        columns,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = cell.font.copy(
            bold=True
        )

    # Data rows
    for row_index, record in enumerate(
        records,
        start=2,
    ):
        for column_index, column_name in enumerate(
            columns,
            start=1,
        ):
            value = record.get(column_name)

            # Convert lists into readable text.
            if isinstance(value, list):
                value = ", ".join(
                    str(item)
                    for item in value
                )

            # Convert dictionaries into JSON text.
            elif isinstance(value, dict):
                value = json.dumps(
                    value,
                    ensure_ascii=False,
                )

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

    # Freeze header row.
    worksheet.freeze_panes = "A2"

    # Add filters.
    if records:
        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

    # Set reasonable column widths.
    for column_index, column_name in enumerate(
        columns,
        start=1,
    ):
        max_length = len(
            str(column_name)
        )

        column_letter = get_column_letter(
            column_index
        )

        for cell in worksheet[column_letter]:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

            if max_length >= 60:
                break

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            60,
        )


def export_to_excel() -> Path:
    """Export all Frontier Atlas datasets into six sheets."""

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # --------------------------------------------------
    # Load datasets
    # --------------------------------------------------

    startups = load_json(
        "startups.json"
    )

    products = load_json(
        "products.json"
    )

    research_papers = load_json(
        "research_papers.json"
    )

    jobs = load_json(
        "jobs.json"
    )

    news = load_json(
        "news.json"
    )

    entity_mapping = load_json(
        "entity_mapping_log.json"
    )

    # --------------------------------------------------
    # Create workbook
    # --------------------------------------------------

    workbook = Workbook()

    # Remove default worksheet safely.
    default_sheet = workbook.active

    if default_sheet is not None:
        workbook.remove(default_sheet)

    # --------------------------------------------------
    # Sheet 1: Startups
    # --------------------------------------------------

    write_sheet(
        workbook=workbook,
        sheet_name="Startups",
        records=startups,
        columns=[
            "schemaVersion",
            "recordType",
            "source_name",
            "source_url",
            "entity_name",
            "employee_count",
            "collected_at",
        ],
    )

    # --------------------------------------------------
    # Sheet 2: Products
    # --------------------------------------------------

    write_sheet(
        workbook=workbook,
        sheet_name="Products",
        records=products,
        columns=[
            "schemaVersion",
            "recordType",
            "source_name",
            "source_url",
            "product_name",
            "startup_name",
            "pricing_model",
            "collected_at",
        ],
    )

    # --------------------------------------------------
    # Sheet 3: Research Papers
    # --------------------------------------------------

    write_sheet(
        workbook=workbook,
        sheet_name="Research Papers",
        records=research_papers,
        columns=[
            "schemaVersion",
            "recordType",
            "title",
            "authors",
            "paper_url",
            "arxiv_id",
            "published_date",
            "github_url",
            "github_stars",
            "github_confidence",
            "summary",
        ],
    )

    # --------------------------------------------------
    # Sheet 4: Jobs
    # --------------------------------------------------

    write_sheet(
        workbook=workbook,
        sheet_name="Jobs",
        records=jobs,
        columns=[
            "schemaVersion",
            "recordType",
            "source_name",
            "source_url",
            "company_name",
            "job_title",
            "location",
            "description",
            "posted_at",
            "collected_at",
        ],
    )

    # --------------------------------------------------
    # Sheet 5: News
    # --------------------------------------------------

    write_sheet(
        workbook=workbook,
        sheet_name="News",
        records=news,
        columns=[
            "schemaVersion",
            "recordType",
            "source_name",
            "source_url",
            "title",
            "description",
            "article_text",
            "published_at",
            "collected_at",
        ],
    )

    # --------------------------------------------------
    # Sheet 6: Entity Mapping Log
    # --------------------------------------------------

    write_sheet(
        workbook=workbook,
        sheet_name="Entity Mapping Log",
        records=entity_mapping,
        columns=[
            "raw_name",
            "canonical_name",
        ],
    )

    # --------------------------------------------------
    # Save workbook
    # --------------------------------------------------

    workbook.save(EXCEL_FILE)

    logger.info(
        "Excel workbook created: %s",
        EXCEL_FILE,
    )

    logger.info(
        "Sheets: %s",
        workbook.sheetnames,
    )

    return EXCEL_FILE


if __name__ == "__main__":

    logging.basicConfig(
        level=logging.INFO,
        format="%(levelname)s: %(message)s",
    )

    output_file = export_to_excel()

    print(
        f"\nEXCEL FILE CREATED: {output_file}"
    )